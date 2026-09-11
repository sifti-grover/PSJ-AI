"""Stage 7b (Reporter Agent's deterministic half): merge insights and export
the final submission-ready report. Called from `09-report.json`.
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def build_docx(project_name: str, clusters: list[dict], out_path: str | Path) -> None:
    """clusters: [{title, differences, similarities, paragraph}, ...]
    Deterministic merge — no LLM call happens in this function.
    """
    doc = Document()
    doc.add_heading(f"{project_name} — Insight Report", level=0)

    for cluster in clusters:
        doc.add_heading(cluster["title"], level=1)
        doc.add_paragraph(cluster["paragraph"])
        doc.add_heading("Similarities", level=2)
        doc.add_paragraph(cluster["similarities"])
        doc.add_heading("Differences", level=2)
        doc.add_paragraph(cluster["differences"])

    doc.save(str(out_path))


def build_prisma_summary(counts: dict[str, int]) -> str:
    """counts: identified/deduped/screened/included -> renders the standard
    PRISMA flow numbers as text; a chart version lives alongside this.
    """
    lines = [f"{stage}: {n}" for stage, n in counts.items()]
    return "\n".join(lines)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-name", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()
    print("Wire this up to pull clusters/insights from Postgres, then call build_docx().")
